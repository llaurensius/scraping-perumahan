#!/usr/bin/env python3
"""Sikumbang Perumahan Subsidi Scraper - Local Web UI"""

import csv
import io
import json
import os
from datetime import datetime
from pathlib import Path

import requests
from flask import Flask, Response, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

SIKUMBANG_BASE = "https://sikumbang.tapera.go.id"
EXPORT_DIR = Path(__file__).parent / "exports"
EXPORT_DIR.mkdir(exist_ok=True)

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
    "Accept": "application/json, text/plain, */*",
    "Referer": "https://sikumbang.tapera.go.id/",
}


def api_get(path, params=None):
    """Proxy GET request to Sikumbang API."""
    url = f"{SIKUMBANG_BASE}{path}"
    try:
        resp = requests.get(url, params=params, headers=HEADERS, timeout=30)
        resp.raise_for_status()
        return resp.json()
    except Exception as e:
        return {"error": str(e)}, 502


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/provinsi")
def get_provinsi():
    data = api_get("/ajax/wilayah/get-provinsi")
    return jsonify(data)


@app.route("/api/kabupaten/<kode>")
def get_kabupaten(kode):
    data = api_get(f"/ajax/wilayah/get-kabupaten/{kode}")
    return jsonify(data)


@app.route("/api/lokasi")
def search_lokasi():
    params = {}
    for key in ("page", "limit", "search", "searchField", "filter", "filterField",
                "ordering", "sortBy"):
        val = request.args.get(key)
        if val is not None:
            params[key] = val
    params.setdefault("page", "1")
    params.setdefault("limit", "100")

    data = api_get("/ajax/lokasi/search", params=params)
    return jsonify(data)


@app.route("/api/lokasi/<id_lokasi>")
def get_lokasi_detail(id_lokasi):
    data = api_get(f"/lokasi-perumahan/{id_lokasi}/json")
    return jsonify(data)


def build_flat_rows(selected_items):
    """Flatten lokasi + tipeRumah into one-row-per-tipe rows."""
    rows = []
    for item in selected_items:
        wilayah = item.get("wilayah", {})
        pengembang = item.get("pengembang", {})
        tipe_list = item.get("tipeRumah", [])

        if not tipe_list:
            tipe_list = [{"nama": "-", "status": "-", "harga": 0,
                          "luasTanah": 0, "luasBangunan": 0,
                          "kamarTidur": 0, "kamarMandi": 0}]

        for tipe in tipe_list:
            kt = tipe.get("kamarTidur", 0)
            km = tipe.get("kamarMandi", 0)
            rows.append({
                "id_lokasi": item.get("idLokasi", ""),
                "nama_perumahan": item.get("namaPerumahan", ""),
                "jenis_perumahan": item.get("jenisPerumahan", ""),
                "provinsi": wilayah.get("provinsi", ""),
                "kab_kota": wilayah.get("kabupaten", ""),
                "kecamatan": wilayah.get("kecamatan", ""),
                "kelurahan": wilayah.get("kelurahan", ""),
                "tipe_rumah": tipe.get("nama", ""),
                "status": tipe.get("status", ""),
                "harga": tipe.get("harga", 0),
                "luas_tanah": tipe.get("luasTanah", 0),
                "luas_bangunan": tipe.get("luasBangunan", 0),
                "kt": kt,
                "km": km,
                "kt_km": f"{kt}/{km}",
                "jumlah_unit": item.get("jumlahUnit", 0),
                "jumlah_unit_komersil": item.get("jumlahUnitKomersil", 0),
                "pengembang": pengembang.get("nama", ""),
                "asosiasi": pengembang.get("asosiasi", ""),
                "koordinat": item.get("koordinatPerumahan", ""),
                "url_detail": f"{SIKUMBANG_BASE}/lokasi-perumahan/{item.get('idLokasi', '')}",
            })
    return rows


CSV_HEADERS = [
    "ID Lokasi", "Nama Perumahan", "Jenis", "Provinsi", "Kab/Kota",
    "Kecamatan", "Kelurahan", "Tipe Rumah", "Status", "Harga",
    "Luas Tanah (m²)", "Luas Bangunan (m²)", "KT/KM",
    "Jumlah Unit (Subsidi)", "Jumlah Unit (Komersil)",
    "Pengembang", "Asosiasi", "Koordinat", "URL Detail",
]


def row_to_list(row):
    return [
        row["id_lokasi"], row["nama_perumahan"], row["jenis_perumahan"],
        row["provinsi"], row["kab_kota"], row["kecamatan"], row["kelurahan"],
        row["tipe_rumah"], row["status"], row["harga"],
        row["luas_tanah"], row["luas_bangunan"], row["kt_km"],
        row["jumlah_unit"], row["jumlah_unit_komersil"],
        row["pengembang"], row["asosiasi"], row["koordinat"], row["url_detail"],
    ]


@app.route("/api/export/csv", methods=["POST"])
def export_csv():
    selected = request.get_json(silent=True) or []
    if not selected:
        return jsonify({"error": "Tidak ada data dipilih"}), 400

    rows = build_flat_rows(selected)
    buf = io.StringIO(newline="")
    buf.write("\ufeff")  # BOM for Excel
    writer = csv.writer(buf, delimiter=",", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(CSV_HEADERS)
    for row in rows:
        writer.writerow(row_to_list(row))

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"sikumbang_{timestamp}.csv"
    filepath = EXPORT_DIR / filename
    filepath.write_text(buf.getvalue(), encoding="utf-8-sig")

    return send_file(
        filepath,
        mimetype="text/csv; charset=utf-8",
        as_attachment=True,
        download_name=filename,
    )


@app.route("/api/export/xlsx", methods=["POST"])
def export_xlsx():
    selected = request.get_json(silent=True) or []
    if not selected:
        return jsonify({"error": "Tidak ada data dipilih"}), 400

    rows = build_flat_rows(selected)

    wb = Workbook()
    ws = wb.active
    ws.title = "Perumahan Subsidi"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    for col_idx, header in enumerate(CSV_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    for row_idx, row in enumerate(rows, 2):
        values = row_to_list(row)
        for col_idx, val in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = thin_border

    harga_col = CSV_HEADERS.index("Harga") + 1
    lt_col = CSV_HEADERS.index("Luas Tanah (m²)") + 1
    lb_col = CSV_HEADERS.index("Luas Bangunan (m²)") + 1
    for row_idx in range(2, len(rows) + 2):
        for col in (harga_col, lt_col, lb_col):
            cell = ws.cell(row=row_idx, column=col)
            if cell.value:
                cell.number_format = "#,##0"

    for col_idx in range(1, len(CSV_HEADERS) + 1):
        max_len = len(CSV_HEADERS[col_idx - 1])
        for row_idx in range(2, len(rows) + 2):
            val = ws.cell(row=row_idx, column=col_idx).value
            if val:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 40)

    ws.freeze_panes = "A2"

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"sikumbang_{timestamp}.xlsx"
    filepath = EXPORT_DIR / filename
    wb.save(filepath)

    return send_file(
        filepath,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )


if __name__ == "__main__":
    print("=" * 50)
    print("  SIKUMBANG - Scraping Perumahan Subsidi")
    print("  Buka browser: http://localhost:8081")
    print("=" * 50)
    app.run(host="0.0.0.0", port=8081, debug=True)
